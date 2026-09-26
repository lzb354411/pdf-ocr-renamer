# -*- coding: utf-8 -*-
"""资料完整性核对测试（阶段五 R-6 / A5-1、A5-2、A5-3）

覆盖：
1. 期望清单生成（进度表读取 + 规则表/模板推导词表）
2. 实收清单扫描（含「未匹配项目」二级目录、前缀剥离、别名归一）
3. 差异三类检出：缺失 / 多余 / 命名不一致（A5-2）
4. 命名方式兼容（by_folder=True / False）
5. 导出 CSV / XLSX 且中文不乱码（A5-3）
6. 边界：空目录、无进度表、项目完全齐全

不依赖 OCR，全部用合成的文件树验证逻辑；
真实进度表用例在缺文件时自动 SKIP。
"""
import os
import sys
import csv
import shutil
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "src"))

from core import inventory as INV

ROOT = os.path.dirname(os.path.abspath(__file__))
SAMPLE_XLSX = os.path.join(ROOT, "2026年业扩项目进度表.xlsx")

failures = []
skips = []


def check(name, cond, detail=""):
    if cond:
        print("  [PASS] {}".format(name))
    else:
        failures.append(name)
        print("  [FAIL] {}  {}".format(name, detail))


def skip(name):
    skips.append(name)
    print("  [SKIP] {}".format(name))


def mkpdf(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("dummy")


# 测试用映射规则（模拟真实规则表的「识别内容 → 映射内容」关系）
RULES = [
    {"识别内容": "配电网工程开工报告", "映射内容": "开工报告"},
    {"识别内容": "配电网工程竣工报告", "映射内容": "竣工报告"},
    {"识别内容": "低压电力电缆试验记录", "映射内容": "0.4kV电缆试验报告"},
    {"识别内容": "监理工作日志", "映射内容": "监理日志"},
]
REQ = ["开工报告", "竣工报告", "0.4kV电缆试验报告", "监理日志"]

print("== 1. 文档类型词表推导 ==")
types_rules = INV.document_types_from_rules(RULES)
check("从规则表推导出 4 种类型", types_rules == ["开工报告", "竣工报告",
                                                "0.4kV电缆试验报告", "监理日志"],
      str(types_rules))
check("空规则返回空表", INV.document_types_from_rules([]) == [])
check("规则含空映射内容时被跳过",
      INV.document_types_from_rules(
          [{"识别内容": "x", "映射内容": ""}]) == [])

# 模板文件名推导（不依赖外部目录）
tmp_tpl = tempfile.mkdtemp()
try:
    for fn in ("1.模板-开工报告.xlsx", "2.模板-竣工报告.xlsx",
               "41.模板-开工报告带签字.xlsx", "15.模板-工作票.xlsx"):
        open(os.path.join(tmp_tpl, fn), "w").close()
    t = INV.document_types_from_templates(tmp_tpl)
    check("模板名去掉序号/模板-/带签字后得到类型",
          "开工报告" in t and "竣工报告" in t and "工作票" in t, str(t))
    check("带签字模板与普通模板去重为同一类型",
          t.count("开工报告") == 1, str(t))
    check("不存在的目录返回空表", INV.document_types_from_templates("无此目录") == [])
finally:
    shutil.rmtree(tmp_tpl, ignore_errors=True)

print("== 2. 别名表（识别内容 ↔ 映射内容 归一） ==")
alias = INV.build_alias_map(RULES)
check("识别内容归一为映射内容",
      alias.get(INV._norm("低压电力电缆试验记录")) == "0.4kV电缆试验报告",
      str(alias.get(INV._norm("低压电力电缆试验记录"))))
check("映射内容自映射", alias.get(INV._norm("开工报告")) == "开工报告")
check("空规则得到空别名表", INV.build_alias_map([]) == {})

print("== 3. 实收清单扫描 ==")
tmp = tempfile.mkdtemp()
try:
    P1 = "项目甲"
    P2 = "项目乙"
    # 项目甲：by_folder=True 命名，齐全
    for t in REQ:
        mkpdf(os.path.join(tmp, P1, P1 + t + ".pdf"))
    # 项目乙：by_folder=False 命名（纯类型名），齐全
    for t in REQ:
        mkpdf(os.path.join(tmp, P2, t + ".pdf"))
    # 未匹配项目/<工程名>/
    mkpdf(os.path.join(tmp, "未匹配项目", "某工程", "某工程开工报告.pdf"))
    # 非 PDF 文件应被忽略
    with open(os.path.join(tmp, P1, "说明.txt"), "w") as f:
        f.write("x")

    actual = INV.scan_actual(tmp, rules=RULES)
    check("扫描到 3 个项目目录（2 正常 + 1 未匹配）", len(actual) == 3,
          str([a["folder"] for a in actual]))
    a1 = [a for a in actual if a["folder"] == P1][0]
    check("PDF 之外的扩展名被忽略", len(a1["files"]) == 4, str(len(a1["files"])))
    check("前缀剥离正确（by_folder=True）",
          set(a1["actual_types"]) == set(REQ), str(a1["actual_types"]))
    a2 = [a for a in actual if a["folder"] == P2][0]
    check("纯类型名正确识别（by_folder=False）",
          set(a2["actual_types"]) == set(REQ), str(a2["actual_types"]))
    a3 = [a for a in actual if a["is_unmatched"]]
    check("未匹配项目被标记", len(a3) == 1 and a3[0]["is_unmatched"],
          str(a3))

    print("== 4. 差异三类检出（A5-2） ==")
    # 项目丙：缺失 + 多余 + 命名不一致 三类齐全
    P3 = "项目丙"
    mkpdf(os.path.join(tmp, P3, P3 + "开工报告.pdf"))          # 应交且正确
    mkpdf(os.path.join(tmp, P3, P3 + "竣工报告.pdf"))          # 应交且正确
    # 缺 0.4kV电缆试验报告、监理日志 → 缺失 2
    mkpdf(os.path.join(tmp, P3, P3 + "环网柜施工记录.pdf"))     # 不在应交 → 多余
    mkpdf(os.path.join(tmp, P3, "别的项目开工报告.pdf"))        # 未以项目丙开头 → 命名不一致

    rep = INV.build_report([P1, P2, P3], REQ, tmp, rules=RULES)
    by = {p["name"]: p for p in rep["projects"]}

    check("项目甲：齐全无问题",
          by[P1]["status"] == INV.KIND_OK and not by[P1]["missing"]
          and not by[P1]["extra"] and not by[P1]["mismatch"],
          str(by[P1]))
    check("项目乙：齐全无问题（纯类型名命名）",
          by[P2]["status"] == INV.KIND_OK and not by[P2]["mismatch"],
          str(by[P2]))

    p3 = by[P3]
    check("检出「缺失」2 项",
          set(p3["missing"]) == {"0.4kV电缆试验报告", "监理日志"},
          str(p3["missing"]))
    check("检出「多余」1 项",
          p3["extra"] == ["环网柜施工记录"], str(p3["extra"]))
    check("检出「命名不一致」1 项",
          p3["mismatch"] == ["别的项目开工报告.pdf"], str(p3["mismatch"]))
    check("项目丙状态为缺失（优先）", p3["status"] == INV.KIND_MISSING,
          p3["status"])

    s = rep["summary"]
    check("汇总：项目数=3", s["projects"] == 3, str(s["projects"]))
    check("汇总：缺失合计=2", s["missing_total"] == 2, str(s["missing_total"]))
    check("汇总：多余合计=1", s["extra_total"] == 1, str(s["extra_total"]))
    check("汇总：命名不一致合计=1", s["mismatch_total"] == 1,
          str(s["mismatch_total"]))
    check("汇总：齐全项目=2", s["ok_projects"] == 2, str(s["ok_projects"]))
    check("汇总：问题项目=1", s["problem_projects"] == 1,
          str(s["problem_projects"]))

    print("== 5. 别名归一：识别内容写法不应被误判 ==")
    P4 = "项目丁"
    # 用「识别内容」写法命名（低压电力电缆试验记录），应交是映射内容
    for t in ("开工报告", "竣工报告", "监理日志"):
        mkpdf(os.path.join(tmp, P4, P4 + t + ".pdf"))
    mkpdf(os.path.join(tmp, P4, P4 + "低压电力电缆试验记录.pdf"))
    rep4 = INV.build_report([P4], REQ, tmp, rules=RULES)
    p4 = rep4["projects"][0]
    check("识别内容写法不被判为缺失",
          "0.4kV电缆试验报告" not in p4["missing"], str(p4["missing"]))
    check("识别内容写法不被判为多余",
          not p4["extra"], str(p4["extra"]))
    check("项目丁齐全", p4["status"] == INV.KIND_OK, str(p4))

    print("== 6. 目录中存在但进度表未列出的项目 ==")
    P5 = "项目戊（未在进度表）"
    mkpdf(os.path.join(tmp, P5, P5 + "开工报告.pdf"))
    rep5 = INV.build_report([P1], REQ, tmp, rules=RULES)
    names = [a["folder"] for a in rep5["actual_only"]]
    check("未列出项目出现在 actual_only 中", P5 in names, str(names))
    check("_前置页面 不计入 actual_only",
          not any(a["folder"] == "_前置页面" for a in rep5["actual_only"]),
          str(names))

    print("== 7. 导出（A5-3） ==")
    rep_all = INV.build_report([P1, P2, P3], REQ, tmp, rules=RULES)
    csv_path = os.path.join(tmp, "out", "报告.csv")
    INV.export_csv(rep_all, csv_path)
    check("CSV 已写出", os.path.exists(csv_path))
    raw = open(csv_path, "rb").read()
    check("CSV 带 UTF-8 BOM（Excel 中文不乱码）", raw.startswith(b"\xef\xbb\xbf"),
          repr(raw[:6]))
    with open(csv_path, encoding="utf-8-sig") as f:
        text = f.read()
    check("CSV 含中文表头", "项目名称" in text and "缺失" in text)
    check("CSV 含具体问题内容", "监理日志" in text)
    # 用 csv 模块解析，确认列数一致
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    check("CSV 可被标准解析器读取（无乱码/无错列）",
          any(len(r) == 3 and r[1] == INV.KIND_MISSING for r in rows),
          str(rows[:3]))

    xlsx_path = os.path.join(tmp, "out", "报告.xlsx")
    got = INV.export_xlsx(rep_all, xlsx_path)
    check("XLSX 已写出", os.path.exists(xlsx_path), got)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(got)
        check("XLSX 含总览与明细两个 sheet",
              "核对总览" in wb.sheetnames and "问题明细" in wb.sheetnames,
              str(wb.sheetnames))
        ws = wb["核对总览"]
        vals = [str(c.value) for row in ws.iter_rows() for c in row
                if c.value is not None]
        check("XLSX 中文内容正确（无乱码）",
              any("项目丙" in v for v in vals), str(vals[:6]))
        wb.close()
    except ImportError:
        skip("openpyxl 不可用，跳过 XLSX 内容校验")

    print("== 8. 边界情形 ==")
    empty = tempfile.mkdtemp()
    try:
        rep_e = INV.build_report([P1], REQ, empty, rules=RULES)
        check("空目录 → 项目全部缺失",
              len(rep_e["projects"][0]["missing"]) == 4, str(rep_e["projects"][0]))
        check("空目录扫描不报错", INV.scan_actual(empty) == [])
    finally:
        shutil.rmtree(empty, ignore_errors=True)
    check("不存在的目录扫描返回空表", INV.scan_actual("无此目录") == [])
    check("空项目列表报告不报错",
          INV.build_report([], REQ, tmp, rules=RULES)["summary"]["projects"] == 0)
    check("缺少必需参数不抛异常（required=None 时用兜底）",
          len(INV._infer_types(None, None, False)) > 0)

    print("== 9. 真实项目进度表（A5-1 前置） ==")
    if not os.path.exists(SAMPLE_XLSX):
        skip("真实进度表用例（未找到 {}）".format(os.path.basename(SAMPLE_XLSX)))
    else:
        names = INV.read_progress_names(SAMPLE_XLSX)
        check("从真实进度表读到 560 条项目名称", len(names) == 560, str(len(names)))
        check("项目名称非空且无重复", len(set(names)) == len(names))
        check("名称包含真实业务词（业扩/维修）",
              any("业扩" in n or "维修" in n for n in names), str(names[:1]))

finally:
    shutil.rmtree(tmp, ignore_errors=True)

print()
if skips:
    print("跳过 {} 项（缺少依赖或样本）".format(len(skips)))
if failures:
    print("结果：{} 项失败 -> {}".format(len(failures), failures))
    sys.exit(1)
print("结果：全部通过")
