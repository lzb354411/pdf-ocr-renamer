# -*- coding: utf-8 -*-
"""资料完整性核对（阶段五 R-6）

职责：把「应交」与「实收」两份清单比对，输出差异报告。

数据来源：
    - **应交（期望）**：项目进度表 Excel（`项目名称` 列）+ 映射规则/模板清单，
      推导每个项目应归档的文档类型。
    - **实收（实际）**：拆分输出目录（`splitter.execute_split` 的产物），
      扫描每个项目子文件夹下实际存在的 PDF。

差异三类（阶段文档 5.3）：
    1. **缺失（missing）**：应交但目录里没有；
    2. **多余（extra）**：目录里有但不在应交清单内（含未知文档）；
    3. **命名不一致（mismatch）**：文档存在，但文件名前缀与项目名不符
       （通常是「未匹配项目」下被人工归位、或拆分时项目名识别有偏差）。

设计约束：
    - 配置驱动：文档类型词表来自映射规则 / 模板文件名，不在代码里硬编码业务词；
    - 纯函数优先：比对逻辑不碰文件系统之外的全局状态，便于测试；
    - 中文不乱码：导出统一用 utf-8-sig（Excel 友好）。
"""

import os
import re
import csv
from datetime import datetime

# ============================================================
# 常量与工具
# ============================================================

# 拆分未匹配时使用的保留目录名（与 splitter 保持一致）
UNMATCHED_DIR = "未匹配项目"
PRE_PAGES_DIR = "_前置页面"

# 差异类型
KIND_MISSING = "缺失"
KIND_EXTRA = "多余"
KIND_MISMATCH = "命名不一致"
KIND_OK = "齐全"

# 默认「应交」文档类型（当无法从规则表/模板推导时的兜底最小集）
# 说明：真实业务中每个项目应交的文档随项目类型不同，
# 因此这里只作为**兜底**，正常应由 build_expected_from_rules 推导。
DEFAULT_REQUIRED = ["开工报告", "竣工报告"]


def _norm(s):
    """归一化文本用于比较：去空白、去标点、转小写"""
    if not s:
        return ""
    s = str(s)
    s = re.sub(r"[\s\u3000]+", "", s)
    s = re.sub(r"[，。、；：？！「」『』《》（）()\[\]【】<>.,;:!?'\"\-_/\\|*#@`~]", "", s)
    return s.lower()


# ============================================================
# 5.1 期望清单（应交）
# ============================================================

def read_progress_names(xlsx_path, max_rows=5000):
    """从项目进度表读取全部「项目名称」

    遍历所有 sheet，在前 10 行内定位「项目名称」表头（兼容表头不在首行），
    返回去重后的名称列表（保持出现顺序）。

    读取失败返回 []（调用方据此提示用户）。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception:
        return []

    names = []
    seen = set()
    try:
        for ws in wb.worksheets:
            try:
                rows = list(ws.iter_rows(values_only=True, max_row=max_rows))
            except Exception:
                continue
            if not rows:
                continue
            col = -1
            header_idx = -1
            for i, row in enumerate(rows[:10]):
                if row is None:
                    continue
                for j, cell in enumerate(row):
                    if cell is not None and str(cell).strip() == "项目名称":
                        col, header_idx = j, i
                        break
                if col >= 0:
                    break
            if col < 0:
                continue
            for row in rows[header_idx + 1:]:
                if row is None or col >= len(row):
                    continue
                v = row[col]
                if v is None:
                    continue
                s = str(v).strip()
                if not s or s in seen:
                    continue
                seen.add(s)
                names.append(s)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return names


def document_types_from_rules(rules):
    """从映射规则推导「文档类型」词表

    取每条规则的 `映射内容`（去重、去空）。这些就是拆分/重命名后
    实际会出现的文档名，因此天然可作「应交资料」的候选词表。
    """
    out = []
    seen = set()
    for r in (rules or []):
        v = str((r or {}).get("映射内容", "") or "").strip()
        if v and v not in seen:
            seen.add(v)
            out.append(v)
    return out


def document_types_from_templates(templates_dir):
    """从模板文件夹名推导「文档类型」词表

    模板文件名形如 `1.模板-开工报告.xlsx` / `41.模板-开工报告带签字.xlsx`，
    去掉「序号.」「模板-」「带签字」后即文档类型。

    用于在用户未提供规则表时，也能得到业务文档词表。
    """
    out = []
    seen = set()
    if not templates_dir or not os.path.isdir(templates_dir):
        return out
    try:
        files = os.listdir(templates_dir)
    except OSError:
        return out
    for fn in files:
        if not fn.lower().endswith((".xlsx", ".xls", ".docx")):
            continue
        name = os.path.splitext(fn)[0]
        # 去掉前导序号：1. / 41. / 20.
        name = re.sub(r"^\d+[\.、]\s*", "", name)
        # 去掉「模板-」「带签字模板-」
        name = re.sub(r"^带签字模板[-—]?", "", name)
        name = re.sub(r"^模板[-—]?", "", name)
        # 去掉尾缀「带签字」
        name = re.sub(r"带签字$", "", name)
        name = name.strip()
        if name and name not in seen:
            seen.add(name)
            out.append(name)
    return out


def build_expected(project_names, required_types, name_key=None):
    """生成「应交清单」

    参数：
        project_names: 项目名称列表（来自进度表）
        required_types: 每个项目应交付的文档类型列表
        name_key: 可选，把项目名称映射为「文件夹名」的函数；
                  None 时原样使用项目名称

    返回 [{"name", "folder", "required": [...], "required_count"}, ...]

    说明：真实业务中不同项目类型应交文档不同（如 10kV 项目才需电缆试验），
    因此本函数只做「统一应交」；如需按类型区分，请传入不同的
    required_types 分组调用后合并（见 build_expected_grouped）。
    """
    req = [t for t in (required_types or []) if t]
    seen = set()
    out = []
    for n in (project_names or []):
        s = str(n).strip()
        if not s or s in seen:
            continue
        seen.add(s)
        folder = name_key(s) if name_key else s
        out.append({
            "name": s,
            "folder": folder,
            "required": list(req),
            "required_count": len(req),
        })
    return out


# ============================================================
# 5.2 实收清单（实际）
# ============================================================

def build_alias_map(rules):
    """构造「文档类型别名表」：归一化名 → 规范名

    背景（真实样本暴露的歧义）：
        同一种资料在目录里可能以**两种名字**出现：
          - **映射内容**（规范名，如「0.4kV电缆试验报告」）—— 拆分/重命名后；
          - **识别内容**（OCR 原标题，如「低压电力电缆试验记录」）—— 若该文件
            是用户手工放置、或未经本工具重命名，就会保留原标题。
        阶段五实测：`低压电力电缆试验记录` 被判为「多余」，
        但它其实就是 `0.4kV电缆试验报告` —— 属**误报**。

    因此这里依据映射规则建立别名：把每条规则的
    `识别内容` 与 `映射内容` 都归一到**同一个**规范名（取 `映射内容`）。

    返回 dict：{归一化键: 规范名}（同时含规范名自身的自映射）。
    """
    alias = {}
    for r in (rules or []):
        if not isinstance(r, dict):
            continue
        canon = str(r.get("映射内容", "") or "").strip()
        src = str(r.get("识别内容", "") or "").strip()
        if not canon:
            continue
        kc = _norm(canon)
        if kc:
            alias[kc] = canon
        ks = _norm(src)
        # 仅当该识别内容尚未被别的规范名占用时才登记，避免多对多时互相覆盖
        if ks and ks not in alias:
            alias[ks] = canon
    return alias


def _canon_type(t, alias):
    """把文档类型规范化为统一名字（无别名表时原样返回）"""
    if not t:
        return t
    if not alias:
        return t
    return alias.get(_norm(t), t)


def scan_actual(root, rules=None):
    """扫描拆分输出目录，统计每个项目实际存在的文档

    参数 root: 拆分输出根目录（其下为各项目子文件夹）
         rules: 可选，映射规则列表；提供时用于建立别名表，
                把「识别内容」形态的文件名归一为「映射内容」规范名
                （避免把同一资料的两种写法误判为「多余」）。

    返回 [{"folder", "name", "path", "files": [{"file", "stem", "type"}],
           "actual_types", "is_unmatched", "is_pre_pages"}, ...]

    识别规则：
        - 一级子目录视为一个项目（`未匹配项目` 视为特殊容器，
          其下二级目录为项目名）；
        - 文件名形如「<项目名><文档类型>.pdf」或「<文档类型>.pdf」；
          文档类型由 `_strip_project_prefix` 结合项目名剥离前缀得到。
    """
    out = []
    if not root or not os.path.isdir(root):
        return out
    alias = build_alias_map(rules)

    def _collect(folder_path, folder_name, is_unmatched):
        try:
            files = [f for f in os.listdir(folder_path)
                     if os.path.isfile(os.path.join(folder_path, f))
                     and f.lower().endswith(".pdf")]
        except OSError:
            files = []
        items = []
        for fn in sorted(files):
            stem = os.path.splitext(fn)[0]
            raw_type = _strip_project_prefix(stem, folder_name)
            items.append({
                "file": fn,
                "stem": stem,
                "type": raw_type,
                "canon_type": _canon_type(raw_type, alias),
            })
        return {
            "folder": folder_name,
            "name": folder_name,
            "path": folder_path,
            "files": items,
            "actual_types": _dedup([it["canon_type"] for it in items
                                    if it["canon_type"]]),
            "is_unmatched": is_unmatched,
            "is_pre_pages": folder_name == PRE_PAGES_DIR,
        }

    for entry in sorted(os.listdir(root)):
        full = os.path.join(root, entry)
        if not os.path.isdir(full):
            continue
        if entry == UNMATCHED_DIR:
            # 「未匹配项目」下再分项目名（以及 _前置页面）
            try:
                subs = sorted(os.listdir(full))
            except OSError:
                subs = []
            for sub in subs:
                subfull = os.path.join(full, sub)
                if os.path.isdir(subfull):
                    out.append(_collect(subfull, sub, True))
        else:
            out.append(_collect(full, entry, False))
    return out


def _dedup(seq):
    seen = set()
    out = []
    for x in seq:
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    return out


def _strip_project_prefix(stem, folder_name):
    """从文件名中剥离项目名前缀，得到文档类型

    拆分命名有两种方式（见 splitter._assign_file_names）：
        - by_folder=True： `<文件夹名><映射内容>.pdf`
        - by_folder=False：`<映射内容>.pdf`
    因此这里先尝试剥离 `folder_name` 前缀；剥不掉则原样作为类型。

    额外兼容「未识别标题」「前置页面_xxx」等工具生成名。
    """
    if not stem:
        return ""
    s = stem
    fn = _norm(folder_name)
    if fn:
        ns = _norm(s)
        if ns.startswith(fn) and len(ns) > len(fn):
            # 归一化长度可能不等（标点被去掉），改用逐字符对齐的近似剥离
            s = _strip_by_normalized(s, folder_name) or s
    if s.startswith("未识别标题"):
        return "未识别标题"
    if s.startswith("前置页面"):
        return "前置页面"
    return s.strip()


def _strip_by_normalized(text, prefix):
    """按归一化后的匹配，剥离 text 开头的 prefix

    归一化会删除标点，导致长度不一致；这里逐字符前进，
    累积的归一化长度达到 len(_norm(prefix)) 即截断。
    """
    target = len(_norm(prefix))
    acc = 0
    for i, ch in enumerate(text):
        if _norm(ch):
            acc += 1
        if acc >= target:
            rest = text[i + 1:]
            return rest.strip()
    return None


# ============================================================
# 5.3 差异比对
# ============================================================

def compare_project(expected_entry, actual_entry, required_types=None,
                    alias=None):
    """比对单个项目，返回差异结果 dict

    返回：
        {
          "name", "folder",
          "required": [...], "actual": [...],
          "required_count", "actual_count",
          "missing": [...],       # 应交未收到
          "extra": [...],         # 收到但不在应交清单
          "mismatch": [...],      # 命名不一致的文件
          "status": KIND_OK / 其他,
        }

    参数 alias: 文档类型别名表（见 build_alias_map）。提供时会把
        实收类型与应交类型都归一为规范名后再比对，避免把同一资料的
        「识别内容」写法（如 低压电力电缆试验记录）误判为「多余」
        （其规范名是 0.4kV电缆试验报告）。

    判定细节：
        - **缺失**：`required` 中归一化后未在 `actual` 出现；
        - **多余**：`actual` 中归一化后未在 `required` 出现；
        - **命名不一致**：文件名**未以项目名（或其归一化形式）开头**，
          说明该文件可能本属别的项目 / 是人工放入的。
    """
    req = [t for t in (required_types if required_types is not None
                       else expected_entry.get("required", [])) if t]
    if actual_entry is None:
        actual_entry = {"files": [], "folder": expected_entry.get("folder", ""),
                        "actual_types": []}

    files = actual_entry.get("files", []) or []
    actual_types = [f.get("type", "") for f in files if f.get("type")]

    req_norm = {}
    for t in req:
        req_norm.setdefault(_norm(_canon_type(t, alias)), t)

    # ---------- 先判「命名不一致」（决定哪些文件算数） ----------
    # 判定思路（避免误报）：
    #   - by_folder=True  → 文件名应以**本项目文件夹名**开头；
    #   - by_folder=False → 文件名应恰好是**某个已知文档类型**。
    #   两者都不满足，说明该文件不属于本项目（或命名被改坏）。
    #
    # 重要：**命名不一致的文件不再计入「多余」**。
    # 否则同一个文件会被重复上报两次（既\"多余\"又\"命名不一致\"），
    # 使汇总数虚高、用户难以判读。二者语义上互斥：
    #   「多余」= 命名正确但清单外的资料；「命名不一致」= 命名本身就有问题。
    folder = expected_entry.get("folder") or actual_entry.get("folder") or ""
    mismatch = []
    ok_files = []
    if folder:
        fn = _norm(folder)
        known = set(req_norm.keys())
        if alias:
            for k, v in alias.items():
                known.add(k)
                known.add(_norm(v))
        for f in files:
            stem = f.get("stem", "")
            if not stem:
                continue
            ns = _norm(stem)
            if (fn and ns.startswith(fn)) or ns in known:
                ok_files.append(f)
            else:
                mismatch.append(f.get("file", ""))
    else:
        ok_files = list(files)

    act_norm = {}
    for f in ok_files:
        canon = _canon_type(f.get("type", ""), alias)
        if canon:
            act_norm.setdefault(_norm(canon), canon)

    missing = [t for k, t in req_norm.items() if k not in act_norm]
    extra = [t for k, t in act_norm.items() if k not in req_norm]

    if missing:
        status = KIND_MISSING
    elif extra or mismatch:
        status = KIND_MISMATCH if mismatch else KIND_EXTRA
    else:
        status = KIND_OK

    # 报告的「实收」= 命名合规且被识别的文档类型（去重）
    reported_actual = _dedup(list(act_norm.values()))

    return {
        "name": expected_entry.get("name", ""),
        "folder": folder,
        "required": req,
        "actual": reported_actual,
        "required_count": len(req),
        "actual_count": len(reported_actual),
        "missing": missing,
        "extra": extra,
        "mismatch": mismatch,
        "status": status,
    }


def _match_actual(expected, actual_list):
    """为期望条目找到对应的实收条目（按名称/文件夹归一化匹配）"""
    if not actual_list:
        return None
    ek = _norm(expected.get("folder") or expected.get("name") or "")
    if not ek:
        return None
    # 精确
    for a in actual_list:
        if _norm(a.get("folder")) == ek:
            return a
    # 包含（文件夹名可能带「（姓名）」等后缀）
    for a in actual_list:
        ak = _norm(a.get("folder"))
        if ak and (ak.startswith(ek) or ek.startswith(ak)):
            return a
    return None


def build_report(project_names, required_types, root,
                 rules=None, templates_dir=None, use_templates=False):
    """一站式生成完整性核对报告

    参数：
        project_names:  项目名称列表（进度表）
        required_types: 应交文档类型（None 时尝试从 rules/模板推导）
        root:           拆分输出根目录
        rules:          映射规则列表（用于推导词表）
        templates_dir:  模板目录（推导词表）
        use_templates:  是否用模板推导词表

    返回 report dict：
        {
          "generated_at", "root",
          "required_types": [...],
          "summary": {"projects", "required_total", "actual_total",
                      "missing_total", "extra_total", "mismatch_total",
                      "ok_projects", "problem_projects"},
          "projects": [ {compare_project 的结果}, ... ],
          "actual_only": [...],   # 目录中存在但进度表未列出的项目
        }
    """
    types = [t for t in (required_types or []) if t]
    if not types:
        types = _infer_types(rules, templates_dir, use_templates)

    alias = build_alias_map(rules)
    expected = build_expected(project_names, types)
    actual_list = scan_actual(root, rules=rules)

    projects = []
    matched_actual_ids = set()
    for exp in expected:
        act = _match_actual(exp, actual_list)
        if act is not None:
            matched_actual_ids.add(id(act))
        projects.append(compare_project(exp, act, types, alias=alias))

    # 只存在于目录、进度表未列出的项目
    actual_only = []
    for a in actual_list:
        if id(a) in matched_actual_ids:
            continue
        if a.get("is_pre_pages"):
            continue
        actual_only.append({
            "folder": a.get("folder"),
            "path": a.get("path"),
            "actual": a.get("actual_types", []),
            "is_unmatched": a.get("is_unmatched", False),
        })

    missing_total = sum(len(p["missing"]) for p in projects)
    extra_total = sum(len(p["extra"]) for p in projects)
    mismatch_total = sum(len(p["mismatch"]) for p in projects)

    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "root": root,
        "required_types": types,
        "summary": {
            "projects": len(projects),
            "required_total": sum(p["required_count"] for p in projects),
            "actual_total": sum(p["actual_count"] for p in projects),
            "missing_total": missing_total,
            "extra_total": extra_total,
            "mismatch_total": mismatch_total,
            "ok_projects": sum(1 for p in projects if p["status"] == KIND_OK),
            "problem_projects": sum(1 for p in projects
                                    if p["status"] != KIND_OK),
            "actual_only_projects": len(actual_only),
        },
        "projects": projects,
        "actual_only": actual_only,
    }


def _infer_types(rules, templates_dir, use_templates):
    """推导应交文档类型（规则表优先，模板兜底）"""
    types = document_types_from_rules(rules)
    if use_templates or not types:
        tpl = document_types_from_templates(templates_dir)
        # 合并但保持顺序：规则表优先
        for t in tpl:
            if t not in types:
                types.append(t)
    return types or list(DEFAULT_REQUIRED)


# ============================================================
# 5.3 导出
# ============================================================

REPORT_HEADERS = ["项目名称", "文件夹", "应交份数", "实收份数", "状态",
                  "缺失", "多余", "命名不一致"]
DETAIL_HEADERS = ["项目名称", "问题类型", "问题内容"]


def export_csv(report, out_path):
    """导出报告为 CSV（utf-8-sig，Excel 打开中文不乱码）

    返回写出路径。
    """
    d = os.path.dirname(os.path.abspath(out_path))
    if d:
        os.makedirs(d, exist_ok=True)
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        s = report.get("summary", {})
        w.writerow(["资料完整性核对报告"])
        w.writerow(["生成时间", report.get("generated_at", "")])
        w.writerow(["核对目录", report.get("root", "")])
        w.writerow(["项目数", s.get("projects", 0),
                    "应交合计", s.get("required_total", 0),
                    "实收合计", s.get("actual_total", 0)])
        w.writerow(["缺失合计", s.get("missing_total", 0),
                    "多余合计", s.get("extra_total", 0),
                    "命名不一致合计", s.get("mismatch_total", 0)])
        w.writerow([])
        w.writerow(REPORT_HEADERS)
        for p in report.get("projects", []):
            w.writerow([p["name"], p["folder"], p["required_count"],
                        p["actual_count"], p["status"],
                        "、".join(p["missing"]),
                        "、".join(p["extra"]),
                        "、".join(p["mismatch"])])
        # 明细表
        w.writerow([])
        w.writerow(DETAIL_HEADERS)
        for p in report.get("projects", []):
            for m in p["missing"]:
                w.writerow([p["name"], KIND_MISSING, m])
            for e in p["extra"]:
                w.writerow([p["name"], KIND_EXTRA, e])
            for mm in p["mismatch"]:
                w.writerow([p["name"], KIND_MISMATCH, mm])
        for a in report.get("actual_only", []):
            w.writerow([a["folder"], KIND_EXTRA,
                        "目录中存在但进度表未列出：" + "、".join(a["actual"])])
    return out_path


def export_xlsx(report, out_path):
    """导出报告为 Excel（openpyxl；不可用时自动回退 CSV）

    返回实际写出的路径。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return export_csv(report, os.path.splitext(out_path)[0] + ".csv")

    wb = Workbook()
    ws = wb.active
    ws.title = "核对总览"
    s = report.get("summary", {})

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    warn_fill = PatternFill("solid", fgColor="FCE4E4")

    ws.append(["资料完整性核对报告"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["生成时间", report.get("generated_at", "")])
    ws.append(["核对目录", report.get("root", "")])
    ws.append(["项目数", s.get("projects", 0),
               "应交合计", s.get("required_total", 0),
               "实收合计", s.get("actual_total", 0)])
    ws.append(["缺失合计", s.get("missing_total", 0),
               "多余合计", s.get("extra_total", 0),
               "命名不一致合计", s.get("mismatch_total", 0)])
    ws.append(["应交文档类型", "、".join(report.get("required_types", []))])
    ws.append([])

    ws.append(REPORT_HEADERS)
    hr = ws.max_row
    for c in ws[hr]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")

    for p in report.get("projects", []):
        ws.append([p["name"], p["folder"], p["required_count"],
                   p["actual_count"], p["status"],
                   "、".join(p["missing"]), "、".join(p["extra"]),
                   "、".join(p["mismatch"])])
        if p["status"] != KIND_OK:
            for c in ws[ws.max_row]:
                c.fill = warn_fill

    # 宽度
    for col, wdt in zip("ABCDEFGH", (40, 40, 10, 10, 12, 30, 30, 30)):
        ws.column_dimensions[col].width = wdt

    # 明细 sheet
    ws2 = wb.create_sheet("问题明细")
    ws2.append(DETAIL_HEADERS)
    for c in ws2[1]:
        c.font = bold
        c.fill = head_fill
    for p in report.get("projects", []):
        for m in p["missing"]:
            ws2.append([p["name"], KIND_MISSING, m])
        for e in p["extra"]:
            ws2.append([p["name"], KIND_EXTRA, e])
        for mm in p["mismatch"]:
            ws2.append([p["name"], KIND_MISMATCH, mm])
    for a in report.get("actual_only", []):
        ws2.append([a["folder"], KIND_EXTRA,
                    "目录中存在但进度表未列出：" + "、".join(a["actual"])])
    for col, wdt in zip("ABC", (40, 12, 60)):
        ws2.column_dimensions[col].width = wdt

    d = os.path.dirname(os.path.abspath(out_path))
    if d:
        os.makedirs(d, exist_ok=True)
    wb.save(out_path)
    return out_path
